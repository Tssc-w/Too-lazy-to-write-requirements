from __future__ import annotations

import copy
import contextvars
import datetime as dt
import html
import hashlib
import hmac
import http.client
from html.parser import HTMLParser
import ipaddress
import json
import mimetypes
import os
import re
import ssl
import socket
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
WEB_DIR = ROOT / "web"
TEMPLATE_PATH = ROOT / "templates" / "做图需求表模版-10.12.xlsx"
PORT = int(os.environ.get("IMAGEREQ_PORT", os.environ.get("PORT", "8766")))
HOST = os.environ.get("IMAGEREQ_HOST", "127.0.0.1").strip() or "127.0.0.1"
ADMIN_TOKEN = os.environ.get("IMAGEREQ_ADMIN_TOKEN", "").strip()
VERCEL_RUNTIME = os.environ.get("VERCEL", "").lower() == "1"
DEFAULT_AI_API_BASE_URL = os.environ.get("OPENAI_API_BASE_URL", "https://api.apimart.ai/v1").strip().rstrip("/") or "https://api.apimart.ai/v1"
OPENAI_VISION_MODEL = os.environ.get("OPENAI_VISION_MODEL", "").strip()
OPENAI_VISION_DETAIL = os.environ.get("OPENAI_VISION_DETAIL", "low").strip().lower()
if OPENAI_VISION_DETAIL not in {"low", "high", "auto"}:
    OPENAI_VISION_DETAIL = "low"
MAX_BODY_BYTES = 12 * 1024 * 1024
MAX_REFERENCE_LINKS = 20
MAX_IMAGES_PER_PRODUCT = 30
MAX_IMAGES_PER_ANALYSIS = 180
ANALYSIS_IMAGE_BATCH_SIZE = 10
ANALYSIS_MAX_WORKERS = 2

_ai_api_key = ""
_ai_api_base_url = DEFAULT_AI_API_BASE_URL
_ai_vision_model = OPENAI_VISION_MODEL
_brightdata_api_token = os.environ.get("BRIGHTDATA_API_TOKEN", "").strip()
_brightdata_zone = os.environ.get("BRIGHTDATA_WEB_UNLOCKER_ZONE", "").strip()
_ai_config_lock = threading.Lock()
_jobs: dict[str, dict] = {}
_jobs_lock = threading.Lock()
_request_ai_config = contextvars.ContextVar("request_ai_config", default=None)


class AIResponseInterruptedError(ValueError):
    """The relay closed an AI response before the full body arrived."""


def text(value, limit=4000):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    return str(value).strip()[:limit]


def safe_url(value):
    value = text(value, 2000)
    try:
        parsed = urllib.parse.urlsplit(value)
        if parsed.scheme not in ("http", "https") or not parsed.netloc or not parsed.hostname or parsed.username or parsed.password:
            return ""
        return urllib.parse.urlunsplit((parsed.scheme.lower(), parsed.netloc, parsed.path or "/", parsed.query, ""))
    except ValueError:
        return ""


def amazon_asin(url):
    try:
        parsed = urllib.parse.urlsplit(url)
        host = (parsed.hostname or "").lower()
        if not (host == "amazon.com" or host.endswith(".amazon.com")):
            return ""
        match = re.search(r"/(?:dp|gp/product)/([A-Z0-9]{10})(?:[/?]|$)", parsed.path, re.I)
        return match.group(1).upper() if match else ""
    except ValueError:
        return ""


def product_identity(url):
    """Stable key for a product link, using ASIN where available and URL otherwise."""
    asin = amazon_asin(url)
    if asin:
        return "amazon:" + asin
    normalized = safe_url(url)
    if not normalized:
        return ""
    parsed = urllib.parse.urlsplit(normalized)
    canonical = urllib.parse.urlunsplit((parsed.scheme, (parsed.hostname or "").lower(), parsed.path.rstrip("/") or "/", parsed.query, ""))
    return "url:" + hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:24]


def _public_http_url(url):
    """Reject local/private destinations before fetching user-supplied URLs."""
    normalized = safe_url(url)
    if not normalized:
        raise ValueError("链接格式无效：仅支持 http:// 或 https:// 网页地址。")
    parsed = urllib.parse.urlsplit(normalized)
    host = (parsed.hostname or "").rstrip(".").lower()
    if host in {"localhost", "localhost.localdomain"} or host.endswith((".localhost", ".local", ".internal", ".test")):
        raise ValueError("出于安全限制，不能读取本机或内网地址。")
    try:
        addresses = {ipaddress.ip_address(host)}
    except ValueError:
        try:
            addresses = {ipaddress.ip_address(info[4][0]) for info in socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)}
        except (OSError, ValueError):
            raise ValueError("无法解析链接域名；请检查地址是否正确或域名是否可公开访问。") from None
    if not addresses or any(not address.is_global for address in addresses):
        raise ValueError("出于安全限制，不能读取指向本机、内网或非公网 IP 的地址。")
    return normalized


class _PublicRedirectHandler(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, request, fp, code, message, headers, new_url):
        _public_http_url(new_url)
        return super().redirect_request(request, fp, code, message, headers, new_url)


def _connect_public_socket(host, port, timeout=socket._GLOBAL_DEFAULT_TIMEOUT, source_address=None):
    try:
        records = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    except OSError:
        raise OSError("无法解析商品页域名。") from None
    if not records or any(not ipaddress.ip_address(record[4][0].split("%", 1)[0]).is_global for record in records):
        raise OSError("出于安全限制，拒绝连接到本机或内网地址。")
    last_error = None
    for family, socktype, proto, _, sockaddr in records:
        connection = socket.socket(family, socktype, proto)
        try:
            if timeout is not socket._GLOBAL_DEFAULT_TIMEOUT:
                connection.settimeout(timeout)
            if source_address:
                connection.bind(source_address)
            connection.connect(sockaddr)
            return connection
        except OSError as exc:
            last_error = exc
            connection.close()
    if last_error:
        raise last_error
    raise OSError("无法连接商品页域名。")


class _PublicHTTPConnection(http.client.HTTPConnection):
    def connect(self):
        self.sock = _connect_public_socket(self.host, self.port, self.timeout, self.source_address)


class _PublicHTTPSConnection(http.client.HTTPSConnection):
    def connect(self):
        sock = _connect_public_socket(self.host, self.port, self.timeout, self.source_address)
        self.sock = self._context.wrap_socket(sock, server_hostname=self.host)


class _PublicHTTPHandler(urllib.request.HTTPHandler):
    def http_open(self, request):
        return self.do_open(_PublicHTTPConnection, request)


class _PublicHTTPSHandler(urllib.request.HTTPSHandler):
    def https_open(self, request):
        return self.do_open(_PublicHTTPSConnection, request, context=self._context, check_hostname=self._check_hostname)


def image_urls(value, base_url=""):
    """Keep unique, absolute HTTP(S) image URLs found in page order."""
    if isinstance(value, str):
        value = [value]
    if isinstance(value, dict):
        value = list(value.values())
    if not isinstance(value, (list, tuple, set)):
        return []
    result, seen = [], set()
    for item in value:
        if isinstance(item, dict):
            item = item.get("hiRes") or item.get("large") or item.get("url") or item.get("image_url") or item.get("src")
        raw_url = html.unescape(text(item, 4000)).replace("\\/", "/").replace("\\u0026", "&")
        if base_url and raw_url:
            raw_url = urllib.parse.urljoin(base_url, raw_url)
        url = safe_url(raw_url)
        if not url or url in seen:
            continue
        host = (urllib.parse.urlsplit(url).hostname or "").lower()
        if host in {"localhost", "127.0.0.1", "::1"} or host.endswith((".localhost", ".local", ".internal")):
            continue
        seen.add(url)
        result.append(url)
        if len(result) >= MAX_IMAGES_PER_PRODUCT:
            break
    return result


def update_job(job_id, **fields):
    with _jobs_lock:
        job = _jobs.get(job_id)
        if job is not None:
            job.update(fields)
            job["updated_at"] = dt.datetime.now().isoformat(timespec="seconds")


def job_snapshot(job_id):
    with _jobs_lock:
        job = _jobs.get(job_id)
        return copy.deepcopy(job) if job else None


class _AmazonHTMLParser(HTMLParser):
    TARGET_IDS = {"productTitle", "feature-bullets", "productDescription", "detailBullets_feature_div", "productDetails_detailBullets_sections1", "prodDetails"}
    IMAGE_REGION_IDS = {"imageBlock_feature_div", "imageBlock", "imgTagWrapperId", "altImages", "landingImage"}
    VOID_TAGS = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.targets = {}
        self.stack = []
        self.meta = {}
        self.image_candidates = []
        self.script_text = []
        self.json_ld = []
        self._script_type = ""
        self._script_buffer = []
        self._in_script = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "meta":
            key = attrs.get("property") or attrs.get("name") or attrs.get("itemprop")
            if key and attrs.get("content"):
                self.meta[key.lower()] = attrs["content"]
        if tag == "script":
            self._in_script = True
            self._script_type = (attrs.get("type") or "").lower()
            self._script_buffer = []
        element_id = attrs.get("id")
        classes = set((attrs.get("class") or "").split())
        is_image_region = element_id in self.IMAGE_REGION_IDS or bool(classes & {"imageBlock", "imgTagWrapperId", "altImages"})
        if element_id in self.TARGET_IDS:
            self.targets.setdefault(element_id, [])
        if tag not in self.VOID_TAGS:
            self.stack.append((tag, element_id if element_id in self.TARGET_IDS else "", is_image_region))
        in_image_region = any(frame[2] for frame in self.stack)
        if in_image_region or element_id == "landingImage" or attrs.get("data-a-dynamic-image"):
            for key in ("data-old-hires", "data-a-hi-res-src", "data-a-hires", "src"):
                if attrs.get(key):
                    self.image_candidates.append(attrs[key])
        for key in ("data-a-dynamic-image",):
            if attrs.get(key):
                value = attrs[key]
                try:
                    parsed = json.loads(html.unescape(value))
                    if isinstance(parsed, dict) and parsed:
                        best = max(parsed.items(), key=lambda pair: (pair[1][0] * pair[1][1]) if isinstance(pair[1], list) and len(pair[1]) >= 2 else 0)
                        if best[0]:
                            self.image_candidates.append(best[0])
                except (TypeError, ValueError, json.JSONDecodeError):
                    self.image_candidates.append(value)
    def handle_endtag(self, tag):
        if tag == "script":
            if "ld+json" in self._script_type:
                try:
                    self.json_ld.append(json.loads("".join(self._script_buffer)))
                except (ValueError, TypeError):
                    pass
            self._in_script = False
            self._script_type = ""
            self._script_buffer = []
        for index in range(len(self.stack) - 1, -1, -1):
            if self.stack[index][0] == tag:
                del self.stack[index:]
                break

    def handle_data(self, data):
        if self._in_script:
            self._script_buffer.append(data)
            if data.strip():
                self.script_text.append(data)
        if data.strip():
            for _, target_id, _ in self.stack:
                if target_id:
                    self.targets[target_id].append(data.strip())


def _clean_html_text(value):
    return re.sub(r"\s+", " ", html.unescape(value or "")).strip()


def _html_section(html_source, pattern, group=1):
    match = re.search(pattern, html_source, re.I | re.S)
    if not match:
        return ""
    return match.group(group)


def _table_specs(html_source):
    specs = []
    for row in re.findall(r"<tr\b[^>]*>(.*?)</tr\s*>", html_source, re.I | re.S):
        cells = re.findall(r"<(?:th|td)\b[^>]*>(.*?)</(?:th|td)\s*>", row, re.I | re.S)
        if len(cells) >= 2:
            label = _clean_html_text(re.sub(r"<[^>]+>", " ", cells[0]))
            value = _clean_html_text(re.sub(r"<[^>]+>", " ", cells[1]))
            if label and value and label.lower() not in ("customer reviews", "best sellers rank"):
                specs.append({"name": label[:160], "value": value[:500]})
    seen, unique = set(), []
    for item in specs:
        key = (item["name"].lower(), item["value"].lower())
        if key not in seen:
            seen.add(key)
            unique.append(item)
    return unique[:60]


def _fetch_product_direct(url, role):
    url = _public_http_url(url)
    asin = amazon_asin(url)
    request = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
    })
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}), _PublicRedirectHandler(), _PublicHTTPHandler(), _PublicHTTPSHandler())
    try:
        with opener.open(request, timeout=28) as response:
            final_url = _public_http_url(response.geturl())
            raw = response.read(6 * 1024 * 1024 + 1)
            if len(raw) > 6 * 1024 * 1024:
                raise ValueError("商品页超过 6 MB，已停止读取。")
            source = raw.decode(response.headers.get_content_charset() or "utf-8", errors="replace")
    except urllib.error.HTTPError as exc:
        raise ValueError(f"页面返回 HTTP {exc.code}。" + ("网站可能要求登录、验证码或限制自动访问。" if exc.code in (401, 403, 429) else "")) from None
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        reason = getattr(exc, "reason", exc)
        raise ValueError("直连网页失败：" + text(reason, 400)) from None
    return parse_product_html(final_url, role, source, "直采")


def fetch_product(url, role, html_source=None, source_method="直采"):
    if html_source is None:
        return _fetch_product_direct(url, role)
    return parse_product_html(url, role, str(html_source), source_method)


def _walk_product_schema(value):
    if isinstance(value, dict):
        kind = value.get("@type") or []
        kinds = [kind] if isinstance(kind, str) else kind if isinstance(kind, list) else []
        if any(str(item).lower().endswith("product") for item in kinds):
            yield value
        for child in value.values():
            yield from _walk_product_schema(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_product_schema(child)


def _jsonld_images(value):
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [value.get("url") or value.get("contentUrl") or value.get("thumbnailUrl") or ""]
    if isinstance(value, list):
        return [url for child in value for url in _jsonld_images(child)]
    return []


def parse_product_html(url, role, source, source_method):
    url = safe_url(url)
    asin = amazon_asin(url) if url else ""
    visible_text = _clean_html_text(re.sub(r"<[^>]+>", " ", source))[:8000]
    if "global adaptive rate limit" in visible_text.lower():
        raise ValueError("Bright Data 返回全局速率限制，请稍后重试。")
    if "robot check" in visible_text.lower() or "validatecaptcha" in source.lower() or "enter the characters you see below" in visible_text.lower() or "captcha" in visible_text.lower()[:1200]:
        raise ValueError("网页返回验证码或机器人校验页，商品详情无法读取。")
    parser = _AmazonHTMLParser()
    parser.feed(source)
    metadata = parser.meta
    schema_products = list(_walk_product_schema(parser.json_ld))
    schema = schema_products[0] if schema_products else {}
    title = _clean_html_text(" ".join(parser.targets.get("productTitle", [])))
    title = title or _clean_html_text(metadata.get("og:title") or metadata.get("twitter:title") or schema.get("name") or _html_section(source, r"<title[^>]*>(.*?)</title>"))
    brand_value = schema.get("brand") or metadata.get("brand") or metadata.get("og:brand") or metadata.get("product:brand")
    brand = _clean_html_text(brand_value.get("name", "") if isinstance(brand_value, dict) else brand_value)
    description = _clean_html_text(" ".join(parser.targets.get("productDescription", [])))
    if not description:
        description = _clean_html_text(metadata.get("description", "") or schema.get("description", ""))
    features = [_clean_html_text(x) for x in parser.targets.get("feature-bullets", [])]
    features = [x for x in dict.fromkeys(features) if len(x) > 12 and x.lower() not in ("› see more product details", "see more product details")][:20]
    specs = []
    for section_id in ("productDetails_detailBullets_sections1", "prodDetails"):
        for match in re.finditer(rf'<table\b[^>]*id=["\']{section_id}["\'][^>]*>.*?</table\s*>', source, re.I | re.S):
            specs.extend(_table_specs(match.group(0)))
    raw_images = list(parser.image_candidates)
    raw_images.extend(_jsonld_images(schema.get("image")))
    color_match = re.search(r"\bcolorImages\s*=\s*", source)
    if color_match:
        object_start = source.find("{", color_match.end())
        if object_start >= 0:
            try:
                color_data, _ = json.JSONDecoder().raw_decode(source[object_start:])
                def collect_media(value):
                    if isinstance(value, dict):
                        for key, child in value.items():
                            if key in ("hiRes", "large", "main", "url") and isinstance(child, str):
                                raw_images.append(child)
                            else:
                                collect_media(child)
                    elif isinstance(value, list):
                        for child in value:
                            collect_media(child)
                collect_media(color_data)
            except (json.JSONDecodeError, RecursionError):
                pass
    imgs = image_urls(raw_images, url)
    main_image = image_urls([metadata.get("og:image", "") or metadata.get("twitter:image", "")], url)
    if main_image and main_image[0] not in imgs:
        imgs.insert(0, main_image[0])
    offers = schema.get("offers") or {}
    if isinstance(offers, list):
        offers = offers[0] if offers else {}
    price = text(offers.get("price") if isinstance(offers, dict) else "", 200)
    price = price or text(metadata.get("product:price:amount") or metadata.get("og:price:amount"), 200)
    currency = text((offers.get("priceCurrency") if isinstance(offers, dict) else "") or metadata.get("product:price:currency") or metadata.get("og:price:currency") or "", 10)
    for pattern in (
        r'<span[^>]+id=["\']priceblock_ourprice["\'][^>]*>(.*?)</span>',
        r'<span[^>]+id=["\']priceblock_dealprice["\'][^>]*>(.*?)</span>',
        r'<span[^>]+class=["\'][^"\']*a-offscreen[^"\']*["\'][^>]*>(.*?)</span>',
    ):
        snippet = _html_section(source, pattern)
        if snippet:
            price = _clean_html_text(re.sub(r"<[^>]+>", " ", snippet))
            if price:
                break
    specifications = schema.get("additionalProperty") or []
    if isinstance(specifications, dict):
        specifications = [specifications]
    if isinstance(specifications, list):
        specs.extend({"name": text(row.get("name"), 160), "value": text(row.get("value"), 500)} for row in specifications if isinstance(row, dict) and row.get("name") and row.get("value"))
    facts = {
        "title": title,
        "brand": brand,
        "price": price,
        "features": features,
        "description": description,
        "specifications": specs,
        "images": imgs,
    }
    field_status = {key: ("已返回" if value else "未返回") for key, value in facts.items()}
    if not (title or imgs or features or description or specs):
        raise ValueError("网页已打开，但没有解析到商品标题、描述或图片。该页面可能不是商品页、需要 JavaScript 渲染、需要登录，或限制自动读取。")
    host = (urllib.parse.urlsplit(url).hostname or "未知网站").lower()
    return {
        "role": role, "platform": host, "asin": asin, "product_id": product_identity(url), "url": url,
        "title": title, "brand": brand, "price": price, "currency": currency,
        "features": features, "description": description, "specifications": specs,
        "images_count": len(imgs), "image_urls": imgs, "image_url": imgs[0] if imgs else "",
        "field_status": field_status, "status": f"{source_method}读取成功" if any(field_status.values()) else "未读取到商品事实",
        "fetch_method": "brightdata" if source_method == "Bright Data" else "direct",
    }


def get_ai_config():
    request_config = _request_ai_config.get()
    if request_config is not None:
        return request_config
    with _ai_config_lock:
        return _ai_api_key, _ai_api_base_url, _ai_vision_model


def submit_with_context(executor, function, *args):
    context = contextvars.copy_context()
    return executor.submit(context.run, function, *args)


def get_brightdata_config():
    with _ai_config_lock:
        return _brightdata_api_token, _brightdata_zone


def fetch_brightdata_product(url, role):
    api_token, zone = get_brightdata_config()
    if not api_token or not zone:
        raise ValueError("Bright Data 回退未配置：请填写 Bright Data API Token 和 Web Unlocker Zone。")
    normalized_url = _public_http_url(url)
    body = json.dumps({"zone": zone, "url": normalized_url, "format": "raw"}).encode("utf-8")
    request = urllib.request.Request("https://api.brightdata.com/request", data=body, method="POST", headers={
        "Authorization": "Bearer " + api_token,
        "Content-Type": "application/json",
        "Accept": "text/html,application/json",
    })
    for attempt in range(2):
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                response_body = response.read(6 * 1024 * 1024 + 1)
            if b"global adaptive rate limit" in response_body.lower() and attempt == 0:
                time.sleep(3)
                continue
            break
        except urllib.error.HTTPError as exc:
            detail = redact_secrets(exc.read(1200).decode("utf-8", errors="replace"), api_token)
            if exc.code in (429, 503) and attempt == 0:
                time.sleep(2)
                continue
            raise ValueError(f"Bright Data API 返回 HTTP {exc.code}：{text(detail, 700)}") from None
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            reason = getattr(exc, "reason", exc)
            raise ValueError("Bright Data API 连接失败或超时：" + redact_secrets(text(reason, 500), api_token)) from None
    if len(response_body) > 6 * 1024 * 1024:
        raise ValueError("Bright Data 返回页面超过 6 MB，已停止读取。")
    source = response_body.decode("utf-8", errors="replace")
    if not source.strip():
        raise ValueError("Bright Data API 返回空页面。")
    try:
        wrapped = json.loads(source)
    except json.JSONDecodeError:
        wrapped = None
    if isinstance(wrapped, dict):
        if wrapped.get("error"):
            raise ValueError("Bright Data 返回错误：" + redact_secrets(text(wrapped.get("error"), 700), api_token))
        source = wrapped.get("body") or wrapped.get("html") or wrapped.get("content") or source
        if not isinstance(source, str):
            source = text(source, 6 * 1024 * 1024)
    return fetch_product(normalized_url, role, html_source=source, source_method="Bright Data")


def normalize_ai_api_base_url(value):
    base_url = text(value, 1200).strip().rstrip("/")
    if not base_url:
        base_url = DEFAULT_AI_API_BASE_URL
    try:
        parsed = urllib.parse.urlsplit(base_url)
        if parsed.scheme not in ("http", "https") or not parsed.hostname or parsed.username or parsed.password:
            raise ValueError
        if parsed.query or parsed.fragment:
            raise ValueError
    except ValueError:
        raise ValueError("AI API Base URL 无效。请填写 http(s):// 开头的地址，不要在地址中附带 Key 或查询参数。") from None
    path = parsed.path.rstrip("/")
    if path.lower().endswith("/chat/completions"):
        endpoint = base_url
    elif path.lower().endswith("/v1"):
        endpoint = base_url + "/chat/completions"
    else:
        endpoint = base_url + "/v1/chat/completions"
    return base_url, endpoint


def fetch_relay_models():
    api_key, api_base_url, _ = get_ai_config()
    if not api_key:
        raise ValueError("请先填写并保存你自己的 AI API Key。")
    _, chat_endpoint = normalize_ai_api_base_url(api_base_url)
    models_endpoint = chat_endpoint.rsplit("/chat/completions", 1)[0] + "/models"
    request = urllib.request.Request(models_endpoint, method="GET", headers={
        "Authorization": "Bearer " + api_key, "Accept": "application/json",
    })
    try:
        with urllib.request.urlopen(request, timeout=20) as response:
            payload = json.loads(response.read(2 * 1024 * 1024).decode("utf-8"))
    except urllib.error.HTTPError as exc:
        raise ValueError(f"AI 模型列表请求返回 HTTP {exc.code}。请确认此 Key 可查询模型列表。") from None
    except (urllib.error.URLError, TimeoutError, OSError, json.JSONDecodeError) as exc:
        reason = getattr(exc, "reason", exc)
        raise ValueError("读取 AI 模型列表失败：" + redact_secrets(text(reason, 400), api_key)) from None
    rows = payload.get("data", payload.get("models", [])) if isinstance(payload, dict) else payload
    if not isinstance(rows, list):
        raise ValueError("AI 服务返回的模型列表格式无法识别。")
    model_ids = []
    for row in rows:
        model_id = text(row.get("id") or row.get("name"), 200) if isinstance(row, dict) else text(row, 200)
        if model_id and model_id not in model_ids:
            model_ids.append(model_id)
        if len(model_ids) >= 500:
            break
    if not model_ids:
        raise ValueError("AI 服务没有返回可选模型；请检查模型列表接口或在设置中手动填写模型 ID。")
    return model_ids


def redact_secrets(value, known_key=""):
    """Remove complete or partially masked API keys from user-visible errors."""
    result = str(value or "")
    if known_key:
        result = result.replace(known_key, "[API Key 已隐藏]")
    return re.sub(r"(?i)\bsk-[A-Za-z0-9_*.-]{4,}", "[API Key 已隐藏]", result)


def extract_sse_chat_content(response_body, api_key, request_id):
    """Collect OpenAI-compatible chat completion content from SSE data events."""
    pieces = []
    for raw_line in response_body.decode("utf-8", errors="replace").lstrip("\ufeff").splitlines():
        line = raw_line.strip()
        if not line or line.startswith(":") or not line.startswith("data:"):
            continue
        event_data = line[5:].strip()
        if not event_data or event_data == "[DONE]":
            continue
        try:
            event = json.loads(event_data)
        except json.JSONDecodeError:
            continue
        if not isinstance(event, dict):
            continue
        if event.get("error"):
            error = event["error"]
            if isinstance(error, dict):
                error = error.get("message") or error.get("type") or "AI 服务返回流式错误"
            raise ValueError("AI 服务流式响应报错：" + redact_secrets(text(error, 500), api_key))
        choices = event.get("choices") or []
        for choice in choices:
            if not isinstance(choice, dict):
                continue
            delta = choice.get("delta") or choice.get("message") or {}
            content = delta.get("content") if isinstance(delta, dict) else None
            if isinstance(content, str):
                pieces.append(content)
            elif isinstance(content, list):
                pieces.extend(part.get("text", "") for part in content if isinstance(part, dict) and isinstance(part.get("text"), str))
            elif isinstance(choice.get("text"), str):
                pieces.append(choice["text"])
    content = "".join(pieces)
    if not content.strip():
        raise ValueError(
            "AI 服务返回了流式响应，但没有可识别的 Chat Completions 文本内容。"
            f"请检查 API 兼容性；请求 ID：{request_id}。"
        )
    return content


def perform_research(job_id, payload):
    try:
        product_url = text(payload.get("product_url"), 2000)
        raw_references = payload.get("reference_urls") or []
        if isinstance(raw_references, str):
            raw_references = raw_references.splitlines()
        links, warnings = [], []
        if product_url:
            links.append((product_url, "我方商品"))
        for raw in raw_references:
            if text(raw):
                links.append((text(raw, 2000), "指定参考"))
        unique, seen = [], set()
        for url, role in links:
            normalized = safe_url(url)
            identity = product_identity(normalized)
            if not normalized:
                warnings.append({"source": text(url, 500), "message": "链接格式无效；仅支持 http:// 或 https:// 网页地址。"})
                continue
            if identity not in seen:
                seen.add(identity)
                unique.append((normalized, role))
            else:
                for index, item in enumerate(unique):
                    if product_identity(item[0]) == identity:
                        if role == "我方商品" or item[1] == "我方商品":
                            unique[index] = (item[0], "我方商品 / 指定参考")
                        break
        items = []
        if unique:
            total = len(unique)
            update_job(job_id, status="正在读取商品页", progress=f"正在尝试直连读取 {total} 个商品页；直连失败时会尝试 Bright Data 回退（如已配置）", completed=0, total=total)
            items = [None] * total
            direct_failures = {}
            with ThreadPoolExecutor(max_workers=min(4, total)) as pool:
                futures = {pool.submit(fetch_product, url, role): (index, url, role) for index, (url, role) in enumerate(unique)}
                completed = 0
                for future in as_completed(futures):
                    index, url, role = futures[future]
                    completed += 1
                    try:
                        items[index] = future.result()
                        items[index]["source_attempts"] = [{"method": "直采", "status": "成功"}]
                    except Exception as exc:
                        direct_error = text(exc, 600)
                        direct_failures[index] = (url, role, direct_error)
                        host = (urllib.parse.urlsplit(url).hostname or "未知网站").lower()
                        items[index] = {
                            "role": role, "platform": host, "asin": amazon_asin(url), "product_id": product_identity(url), "url": url,
                            "title": "", "brand": "", "price": "", "currency": "", "features": [],
                            "description": "", "specifications": [], "images_count": 0, "image_urls": [],
                            "image_url": "", "field_status": {}, "status": "直采失败：" + direct_error,
                            "fetch_method": "failed", "source_attempts": [{"method": "直采", "status": direct_error}],
                        }
                    update_job(job_id, progress=f"已处理 {completed}/{total} 个商品页", completed=completed, total=total)
            if direct_failures:
                api_token, zone = get_brightdata_config()
                if api_token and zone:
                    update_job(job_id, status="直采失败，Bright Data 正在回退", progress=f"直采成功 {total - len(direct_failures)}/{total}；正在通过 Bright Data 重试 {len(direct_failures)} 条失败链接", completed=total - len(direct_failures), total=total)
                    with ThreadPoolExecutor(max_workers=min(2, len(direct_failures))) as pool:
                        futures = {
                            pool.submit(fetch_brightdata_product, url, role): (index, url, role, direct_error)
                            for index, (url, role, direct_error) in direct_failures.items()
                        }
                        for future in as_completed(futures):
                            index, url, role, direct_error = futures[future]
                            try:
                                item = future.result()
                                item["source_attempts"] = [
                                    {"method": "直采", "status": direct_error},
                                    {"method": "Bright Data", "status": "成功"},
                                ]
                                item["status"] = f"Bright Data 回退读取成功；直采失败原因：{direct_error}"
                                items[index] = item
                            except Exception as exc:
                                brightdata_error = redact_secrets(text(exc, 700), api_token)
                                failure = f"直采失败：{direct_error}；Bright Data 回退失败：{brightdata_error}"
                                items[index]["status"] = failure
                                items[index]["source_attempts"].append({"method": "Bright Data", "status": brightdata_error})
                                warnings.append({"source": url, "message": failure})
                            update_job(job_id, progress=f"直采失败链接已由 Bright Data 处理；当前完成 {sum(1 for item in items if item and item.get('fetch_method') != 'failed')}/{total}", completed=sum(1 for item in items if item and item.get("fetch_method") != "failed"), total=total)
                else:
                    for index, (url, role, direct_error) in direct_failures.items():
                        failure = f"直采失败：{direct_error}；Bright Data 未配置，无法回退。请在设置中填写 Bright Data API Token 和 Web Unlocker Zone。"
                        items[index]["status"] = failure
                        warnings.append({"source": url, "message": failure})
            items = [item for item in items if item]
        else:
            if not warnings:
                warnings.append({"source": "商品链接", "message": "没有有效的商品链接；本流程读取已提供的商品页，不自动搜索竞品。"})
        name = text(payload.get("product_name"), 300) or next((i["title"] for i in items if i.get("title")), "")
        summary = {
            "product_name": name, "sku": text(payload.get("sku"), 100), "items": items,
            "warnings": warnings, "searched_at": dt.datetime.now().isoformat(timespec="seconds"),
            "collection_method": "支持任意公开商品网页；优先直采，失败后对该链接尝试 Bright Data 回退",
        }
        update_job(job_id, status="采集完成" if items else "未采集到商品", progress="商品页读取结束。未返回字段以缺失状态显示。", result=summary, done=True, completed=len(items), total=len(items))
    except Exception as exc:
        update_job(job_id, status="failed", error="采集任务中断：" + text(exc, 1000), done=True)


def call_openai_chat(messages, max_tokens=16000, timeout=180):
    api_key, api_base_url, vision_model = get_ai_config()
    if not api_key:
        raise ValueError("请先在「AI 设置」中填写并保存你自己的 API Key。")
    request_payload = {
        "messages": messages,
        "response_format": {"type": "json_object"}, "temperature": 0.2, "max_tokens": max_tokens,
    }
    # Leave model unset when the user delegates model selection to the relay.
    if vision_model:
        request_payload["model"] = vision_model
    body = json.dumps(request_payload).encode("utf-8")
    _, api_endpoint = normalize_ai_api_base_url(api_base_url)
    request = urllib.request.Request(api_endpoint, data=body, method="POST", headers={
        "Authorization": "Bearer " + api_key, "Content-Type": "application/json", "Accept": "application/json",
    })
    for attempt in range(3):
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                response_status = response.status
                response_headers = response.headers
                response_body = response.read()
            content_type = response_headers.get("Content-Type", "未提供")
            request_id = next((response_headers.get(name) for name in ("x-request-id", "request-id", "x-correlation-id", "cf-ray") if response_headers.get(name)), "未提供")
            response_content = None
            if "text/event-stream" in content_type.lower() or response_body.lstrip().startswith(b"data:"):
                response_content = extract_sse_chat_content(response_body, api_key, request_id)
            else:
                try:
                    payload = json.loads(response_body.decode("utf-8-sig"))
                except (UnicodeDecodeError, json.JSONDecodeError):
                    raise ValueError(
                        f"AI 服务返回 HTTP {response_status}，响应格式无法识别。"
                        f"Content-Type：{content_type}；响应长度：{len(response_body)} 字节；请求 ID：{request_id}。"
                    ) from None
            break
        except urllib.error.HTTPError as exc:
            detail = exc.read(4000).decode("utf-8", errors="replace")
            content_type = exc.headers.get("Content-Type", "未提供") if exc.headers else "未提供"
            request_id = next((exc.headers.get(name) for name in ("x-request-id", "request-id", "x-correlation-id", "cf-ray") if exc.headers and exc.headers.get(name)), "未提供")
            if exc.code == 401:
                # The upstream response may echo a partially masked key; never show it.
                raise ValueError(f"AI API 返回 HTTP 401：API Key 无效、已撤销，或服务商不接受此 Key。请核对 Key 和 API 地址。Content-Type：{content_type}；请求 ID：{request_id}。") from None
            try:
                parsed = json.loads(detail)
                if isinstance(parsed, dict):
                    error_value = parsed.get("error") or parsed.get("message") or parsed
                    if isinstance(error_value, dict):
                        error_value = error_value.get("message") or error_value
                    detail = text(error_value, 1200)
                else:
                    detail = text(parsed, 1200)
            except json.JSONDecodeError:
                detail = detail[:1200]
            detail = redact_secrets(detail, api_key)
            raise ValueError(f"AI API 返回 HTTP {exc.code}：{detail}（Content-Type：{content_type}；请求 ID：{request_id}）") from None
        except (http.client.IncompleteRead, http.client.RemoteDisconnected) as exc:
            if attempt < 2:
                time.sleep(attempt + 1)
                continue
            raise AIResponseInterruptedError(
                f"AI 服务响应传输中断，连续重试 3 次仍未读完整（{text(exc, 300)}）。"
            ) from None
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            reason = getattr(exc, "reason", exc)
            is_tls_eof = (
                isinstance(reason, (ssl.SSLEOFError, http.client.RemoteDisconnected))
                or "UNEXPECTED_EOF_WHILE_READING" in str(reason)
            )
            if is_tls_eof:
                if attempt < 2:
                    time.sleep(attempt + 1)
                    continue
                raise AIResponseInterruptedError(
                    "AI 服务 TLS 连接提前关闭，连续重试 3 次仍失败。"
                ) from None
            raise ValueError("AI API 连接失败或超时：" + text(reason, 800)) from None
    if response_content is None:
        try:
            response_content = payload["choices"][0]["message"]["content"]
        except (KeyError, IndexError, TypeError):
            raise ValueError(f"AI 服务响应为 JSON，但缺少 choices[0].message.content；请求 ID：{request_id}。") from None
    if not isinstance(response_content, str) or not response_content.strip():
        raise ValueError(f"AI 服务响应为 JSON，但模型内容为空；请求 ID：{request_id}。")
    try:
        result = json.loads(response_content)
    except json.JSONDecodeError:
        preview = redact_secrets(response_content, api_key).replace("\r", " ").replace("\n", " ")[:500]
        raise ValueError(f"AI 模型内容不是有效 JSON（内容长度：{len(response_content)} 字符；请求 ID：{request_id}）。内容摘要：{preview}") from None
    if not isinstance(result, dict):
        raise ValueError("OpenAI 返回内容结构无效，无法生成套图分析。")
    return result


def call_openai_vision(references, our_product, same_product, user_requirements="", source_image_start=1):
    api_key = get_ai_config()[0]
    if not api_key:
        raise ValueError("请先在「AI 设置」中填写并保存你自己的 API Key。")
    images_payload = []
    reference_facts = []
    for reference_index, reference in enumerate(references, 1):
        reference_facts.append({
            "reference_index": reference_index,
            **{key: reference.get(key) for key in ("title", "asin", "url", "brand", "price", "features", "description", "specifications") if reference.get(key)},
        })
        for offset, url in enumerate(reference.get("image_urls") or []):
            source_image_index = source_image_start + offset
            if source_image_index > MAX_IMAGES_PER_PRODUCT:
                break
            images_payload.append({
                "index": len(images_payload) + 1, "reference_index": reference_index,
                "source_image_index": source_image_index, "asin": reference.get("asin", ""),
                "title": reference.get("title", ""), "source_url": reference.get("url", ""), "url": url,
            })
    if not images_payload:
        raise ValueError("所选参考商品没有可读取的图片 URL。页面字段和来源状态已保留。")
    if len(images_payload) > MAX_IMAGES_PER_ANALYSIS:
        raise ValueError(f"当前选择共 {len(images_payload)} 张图片，单次联合分析最多支持 {MAX_IMAGES_PER_ANALYSIS} 张。请取消部分链接后重试。")
    allowed_types = ["促销主图", "效果展示图", "产品展示/场景图", "卖点图", "色卡/变体图", "细节特写图", "教程/步骤图", "信任背书图", "对比图", "规格参数图"]
    our_facts = our_product or {}
    system = (
        "你是跨平台电商产品图片分析与视觉需求顾问。严格区分商品页事实和图片观察。"
        "尽量遵守用户补充的设计要求，但用户要求不能作为商品事实或授权你补造商品属性。"
        "不得补造规格、材料、认证、折扣、性能、排名、评论或销量。短文案只能依据商品页已返回卖点。"
        "你只需为后续跨图片归纳收集分类标签和极简观察，不要逐张写分析、判断依据、文案或生图指令。"
        "按输入顺序内部分类；产品展示图和场景图统一归入产品展示/场景图。其余类型只能从以下列表中选：" + "、".join(allowed_types) + "。"
        "每张图片只返回 index、type、visual_observation；观察限于可见元素，不写推测。"
        "以 JSON 对象返回字段 product_summary_zh, core_benefits(array), differentiation_zh, source_basis_zh, overall_structure(object: sequence, style, composition, color, material_presentation, copy_strategy), image_analyses(array: index,type,visual_observation)。"
        "商品页事实与图片观察必须区分，不得补造规格、材料、认证、折扣、性能、排名、评论或销量。"
        "当我方商品事实缺失时，product_summary_zh 只概括参考商品，differentiation_zh 说明缺少我方商品信息。"
        "当 same_product 为 true 时，differentiation_zh 明确说明页面证据未支持我方与参考商品之间的差异化结论。"
    )
    context = {
        "same_product": bool(same_product),
        "user_design_requirements": text(user_requirements, 2000),
        "our_product_page_facts": {key: our_facts.get(key) for key in ("title", "brand", "price", "features", "description", "specifications") if our_facts.get(key)},
        "reference_product": reference_facts[0],
        "reference_images_in_page_order": [{key: image[key] for key in ("index", "source_image_index", "asin", "title", "source_url", "url")} for image in images_payload],
    }
    content = [{"type": "text", "text": "请识别此参考商品套图中的类型与可见共性，返回中的分类和观察须遵守系统规则。\n" + json.dumps(context, ensure_ascii=False)}]
    for item in images_payload:
        content.append({"type": "text", "text": f"参考商品的第 {item['source_image_index']} 张；图片序号 {item['index']}；来源 {item['source_url']}"})
        content.append({"type": "image_url", "image_url": {"url": item["url"], "detail": OPENAI_VISION_DETAIL}})
    result = call_openai_chat([
        {"role": "system", "content": system}, {"role": "user", "content": content},
    ], max_tokens=8000, timeout=180)
    if isinstance(result.get("core_benefits"), str):
        result["core_benefits"] = [part.strip() for part in re.split(r"[;；\n]", result["core_benefits"]) if part.strip()]
    elif not isinstance(result.get("core_benefits"), list):
        result["core_benefits"] = []
    has_our_facts = bool(our_facts.get("title") or our_facts.get("features") or our_facts.get("description") or our_facts.get("specifications"))
    result["reference_products"] = [{key: reference.get(key) for key in ("title", "asin", "url", "image_url", "brand", "price")} for reference in references]
    result["reference_product"] = result["reference_products"][0]
    result["our_product"] = {key: our_facts.get(key) for key in ("title", "asin", "url", "brand", "price")} if has_our_facts else None
    result["same_product"] = bool(same_product)
    result["model"] = get_ai_config()[2] or "服务默认模型"
    result["images"] = []
    generated = {int(row.get("index", 0)): row for row in result.get("image_analyses", []) if isinstance(row, dict)}
    allowed_types = {"促销主图", "效果展示图", "产品展示/场景图", "卖点图", "色卡/变体图", "细节特写图", "教程/步骤图", "信任背书图", "对比图", "规格参数图"}
    for item in images_payload:
        row = generated.get(item["index"], {})
        image_type = text(row.get("type"), 100)
        if image_type in {"产品展示图", "场景图"}:
            image_type = "产品展示/场景图"
        if image_type not in allowed_types:
            image_type = "产品展示/场景图"
        result["images"].append({
            "index": item["index"], "source_image_index": item["source_image_index"],
            "reference_index": item["reference_index"], "reference_title": item["title"],
            "image_url": item["url"],
            "type": image_type, "visual_observation": row.get("visual_observation", ""),
            "source_url": item["source_url"], "asin": item["asin"],
        })
    result.pop("image_analyses", None)
    return result


def call_openai_vision_resilient(references, our_product, same_product, user_requirements="", source_image_start=1):
    """Retry an interrupted vision batch with fewer images until it succeeds or reaches one image."""
    try:
        return call_openai_vision(references, our_product, same_product, user_requirements, source_image_start)
    except AIResponseInterruptedError:
        if len(references) != 1:
            raise
        reference = references[0]
        image_urls = reference.get("image_urls") or []
        if len(image_urls) <= 1:
            raise
        midpoint = len(image_urls) // 2
        left_reference = dict(reference, image_urls=image_urls[:midpoint])
        right_reference = dict(reference, image_urls=image_urls[midpoint:])
        left = call_openai_vision_resilient(
            [left_reference], our_product, same_product, user_requirements, source_image_start
        )
        right = call_openai_vision_resilient(
            [right_reference], our_product, same_product, user_requirements, source_image_start + midpoint
        )
        merged = dict(left)
        merged["images"] = (left.get("images") or []) + (right.get("images") or [])
        merged["core_benefits"] = list(dict.fromkeys(
            (left.get("core_benefits") or []) + (right.get("core_benefits") or [])
        ))
        if right.get("overall_structure"):
            merged["overall_structure_by_batch"] = [
                item for item in (left.get("overall_structure"), right.get("overall_structure")) if item
            ]
        return merged


def summarize_reference_analyses(references, analyses, our_product, same_product, user_requirements=""):
    display_types = ["产品展示/场景图", "促销主图", "效果展示图", "卖点图", "色卡/变体图", "细节特写图", "教程/步骤图", "信任背书图", "对比图", "规格参数图"]
    def normalized_type(value):
        value = text(value, 100)
        if value in {"产品展示图", "场景图", "产品展示/场景图", "产品展示图/场景图", "产品展示与场景图"}:
            return "产品展示/场景图"
        return value

    grouped = {kind: {"count": 0, "observations": [], "sources": [], "example_images": [], "example_urls": set()} for kind in display_types}
    source_rows = []
    total_images = 0
    for index, (reference, analysis) in enumerate(zip(references, analyses), 1):
        images = analysis.get("images", [])
        total_images += len(images)
        type_counts = {}
        per_reference_samples = {}
        for image in images:
            image_type = normalized_type(image.get("type") or "产品展示/场景图")
            if image_type not in grouped:
                image_type = "产品展示/场景图"
            type_counts[image_type] = type_counts.get(image_type, 0) + 1
            entry = grouped[image_type]
            entry["count"] += 1
            if reference.get("url") and reference["url"] not in entry["sources"]:
                entry["sources"].append(reference["url"])
            example_url = text(image.get("image_url"), 2000)
            if example_url and example_url not in entry["example_urls"] and len(entry["example_images"]) < 4:
                entry["example_urls"].add(example_url)
                entry["example_images"].append({
                    "image_url": example_url,
                    "source_url": text(image.get("source_url") or reference.get("url"), 2000),
                    "reference_title": text(image.get("reference_title") or reference.get("title") or reference.get("asin"), 300),
                    "sequence_position": image.get("source_image_index"),
                    "visual_observation": text(image.get("visual_observation"), 220),
                })
            samples = per_reference_samples.setdefault(image_type, 0)
            if samples < 2 and image.get("visual_observation"):
                entry["observations"].append({
                    "reference": reference.get("title") or reference.get("asin") or f"参考链接 {index}",
                    "visual_observation": text(image.get("visual_observation"), 220),
                    "sequence_position": image.get("source_image_index"),
                })
                per_reference_samples[image_type] += 1
        source_rows.append({
            "reference_index": index,
            "product_page_facts": {key: reference.get(key) for key in ("title", "asin", "url", "brand", "price", "features", "description", "specifications") if reference.get(key)},
            "product_summary": analysis.get("product_summary_zh", ""),
            "core_benefits": analysis.get("core_benefits", []),
            "image_sequence_structures": analysis.get("overall_structure_by_batch") or [analysis.get("overall_structure", {})],
            "image_count": len(images), "image_type_counts": type_counts,
        })
    visual_groups = [{
        "type": kind, "image_count": info["count"],
        "reference_links": info["sources"], "representative_visual_observations": info["observations"][:12],
        "example_images": info["example_images"],
    } for kind, info in grouped.items() if info["count"]]
    summary_group = next((group for group in visual_groups if group["type"] == "产品展示/场景图"), None)
    # The merged product-display / lifestyle category gets both its four shared
    # findings and two actionable design proposals. Keep the proposals type-level;
    # never reintroduce per-image analysis.
    suggestion_groups = visual_groups
    system = (
        "你是跨平台电商套图研究顾问。只根据给出的商品页事实、图片分类和画面观察归纳，不补造规格、材料、认证、排名、销量或图片中不可见的属性。"
        "产品展示图与场景图已合并为同一个类型。summary_points 必须恰好给出四条该合并类型图片的共同点总结；不得逐张叙述、不得引用单张图序。若输入没有该类型，则总结其余套图共同规律。"
        "image_groups 覆盖输入中出现的每一种图片类型，包括合并后的产品展示/场景图。产品展示/场景图除四条共性总结外，还必须给出恰好两个不同、可执行的类型级设计建议；其余每种出现的类型也必须给出恰好两个不同建议。不得逐张分析或针对单张图写建议。建议结合该类型观察到的共性，不得重复换词。"
        "每条建议都必须有画面文案/卖点文字 copy_zh，写成可直接排版在图片上的短句；若商品事实不足以支持文案或我方定制指令，对应字段留空，不得编造卖点。英文文案 copy_en 可提供自然英文版本。"
        "建议方向应适用于我方商品，不能照搬竞品品牌、商标或原图文字。prompt_zh 和 prompt_en 描述画面构图、背景、光线、材质表现、文字位置及应保留的产品特征。"
        "overall_structure 总结套图整体的顺序、视觉风格、构图、色彩、材质和文案安排。以 JSON 返回，不要 Markdown。"
        "字段为 product_summary_zh, core_benefits(array), differentiation_zh, source_basis_zh, summary_points(array of exactly 4 strings), overall_structure(object: sequence, style, composition, color, material_presentation, copy_strategy), image_groups(array of {type,image_count,shared_patterns,suggestions(array of exactly 2 {title,direction,buyer_focus_zh,copy_zh,copy_en,prompt_zh,prompt_en})})."
        "如我方商品事实缺失，只提供参考商品摘要；copy_zh, copy_en, prompt_zh, prompt_en 留空。same_product 为 true 时，明确页面证据不支持我方与参考商品有差异。"
    )
    context = {
        "same_product": bool(same_product),
        "our_product_page_facts": {key: our_product.get(key) for key in ("title", "asin", "url", "brand", "price", "features", "description", "specifications") if our_product and our_product.get(key)},
        "user_design_requirements": text(user_requirements, 2000),
        "reference_analyses": source_rows,
        "images_grouped_by_type": [{key: value for key, value in group.items() if key != "example_images"} for group in visual_groups],
    }
    result = call_openai_chat([
        {"role": "system", "content": system},
        {"role": "user", "content": "请从已汇总的多张商品图中提取重复出现的共性，输出四条套图总结，并为每个出现的图片类型（包括产品展示/场景图合并类别）各给两项不同的类型级方案；不要逐张分析。\n" + json.dumps(context, ensure_ascii=False)},
    ], max_tokens=8000, timeout=180)
    result["summary_points"] = [text(point, 300) for point in result.get("summary_points", []) if text(point)] if isinstance(result.get("summary_points"), list) else []
    result["summary_points"] = result["summary_points"][:4]
    while len(result["summary_points"]) < 4:
        result["summary_points"].append(f"第 {len(result['summary_points']) + 1} 条共性结论未能可靠生成，请重新分析或人工补充。")
    result["core_benefits"] = result.get("core_benefits") if isinstance(result.get("core_benefits"), list) else []
    has_our_facts = bool(our_product and any(our_product.get(key) for key in ("title", "features", "description", "specifications")))
    generated_groups = {normalized_type(group.get("type")): group for group in result.get("image_groups", []) if isinstance(group, dict)} if isinstance(result.get("image_groups"), list) else {}
    result["image_groups"] = []
    for visual_group in suggestion_groups:
        generated_group = generated_groups.get(visual_group["type"], {})
        suggestions = generated_group.get("suggestions", []) if isinstance(generated_group.get("suggestions"), list) else []
        normalized_suggestions = []
        seen_directions = set()
        for suggestion in suggestions:
            if not isinstance(suggestion, dict):
                continue
            direction = text(suggestion.get("direction"), 700)
            if not direction or direction.casefold() in seen_directions:
                continue
            seen_directions.add(direction.casefold())
            normalized_suggestions.append({
                "title": text(suggestion.get("title"), 120),
                "direction": direction,
                "buyer_focus_zh": text(suggestion.get("buyer_focus_zh"), 400),
                "copy_zh": text(suggestion.get("copy_zh"), 180) if has_our_facts else "",
                "copy_en": text(suggestion.get("copy_en"), 180) if has_our_facts else "",
                "prompt_zh": text(suggestion.get("prompt_zh"), 3000) if has_our_facts else "",
                "prompt_en": text(suggestion.get("prompt_en"), 3000) if has_our_facts else "",
            })
            if len(normalized_suggestions) == 2:
                break
        while len(normalized_suggestions) < 2:
            normalized_suggestions.append({
                "title": f"待补充方案 {len(normalized_suggestions) + 1}", "direction": "",
                "buyer_focus_zh": "", "copy_zh": "", "copy_en": "", "prompt_zh": "", "prompt_en": "",
            })
        result["image_groups"].append({
            "type": visual_group["type"], "image_count": visual_group["image_count"],
            "shared_patterns": text(generated_group.get("shared_patterns"), 700),
            "reference_sources": visual_group["reference_links"],
            "example_images": visual_group["example_images"],
            "suggestions": normalized_suggestions,
        })
    result["image_count"] = total_images
    result["summary_scope"] = "产品展示/场景图共性" if summary_group else "所选套图共性"
    result["summary_image_count"] = summary_group["image_count"] if summary_group else total_images
    result["reference_products"] = [{key: reference.get(key) for key in ("title", "asin", "url", "image_url", "brand", "price")} for reference in references]
    result["reference_product"] = result["reference_products"][0] if result["reference_products"] else {}
    result["our_product"] = {key: our_product.get(key) for key in ("title", "asin", "url", "brand", "price")} if has_our_facts else None
    result["same_product"] = bool(same_product)
    return result


def perform_analysis(job_id, source_job, payload, ai_config):
    context_token = _request_ai_config.set(ai_config)
    try:
        research = (source_job or {}).get("result") or {}
        items = research.get("items") or []
        selected_product_ids = payload.get("selected_product_ids") or payload.get("selected_asins") or ([payload.get("selected_asin")] if payload.get("selected_asin") else [])
        if isinstance(selected_product_ids, str):
            selected_product_ids = [part.strip() for part in selected_product_ids.split(",") if part.strip()]
        selected_product_ids = list(dict.fromkeys(text(product_id, 100) for product_id in selected_product_ids if text(product_id, 100)))
        if not selected_product_ids:
            raise ValueError("请至少选择一条参考商品链接。")
        references = []
        for selected_product_id in selected_product_ids:
            reference = next((item for item in items if (item.get("product_id") or item.get("asin")) == selected_product_id and item.get("role") != "我方商品"), None)
            if not reference:
                reference = next((item for item in items if (item.get("product_id") or item.get("asin")) == selected_product_id), None)
            if not reference:
                raise ValueError(f"找不到所选参考商品 {selected_product_id}，请重新采集链接。")
            if not reference.get("image_urls"):
                raise ValueError(f"所选参考商品 {selected_product_id} 没有可用图片，请取消该链接后重试。")
            references.append(reference)
        image_count = sum(min(len(reference.get("image_urls") or []), MAX_IMAGES_PER_PRODUCT) for reference in references)
        if image_count > MAX_IMAGES_PER_ANALYSIS:
            raise ValueError(f"当前选择共 {image_count} 张图片，单次联合分析最多支持 {MAX_IMAGES_PER_ANALYSIS} 张。请取消部分链接后重试。")
        same_product = bool(payload.get("same_product"))
        our_product_id = text(payload.get("our_product_id"), 100) or text(payload.get("our_asin"), 20).upper()
        our_product = (next((item for item in references if (item.get("product_id") or item.get("asin")) == our_product_id), references[0])) if same_product else next((item for item in items if (item.get("product_id") or item.get("asin")) == our_product_id and str(item.get("role", "")).startswith("我方商品")), None)
        update_job(job_id, status="AI 正在并行分析", progress=f"正在并行分析 {len(references)} 条商品链接、{image_count} 张图片（图像细节模式：{OPENAI_VISION_DETAIL}）", completed=0, total=image_count)
        batches = []
        for reference_index, reference in enumerate(references):
            image_list = (reference.get("image_urls") or [])[:MAX_IMAGES_PER_PRODUCT]
            for start in range(0, len(image_list), ANALYSIS_IMAGE_BATCH_SIZE):
                batch_reference = dict(reference)
                batch_reference["image_urls"] = image_list[start:start + ANALYSIS_IMAGE_BATCH_SIZE]
                batches.append((reference_index, reference, start, batch_reference))
        per_reference_batches = [[] for _ in references]
        completed_batches = 0
        completed_images = 0
        with ThreadPoolExecutor(max_workers=min(ANALYSIS_MAX_WORKERS, len(batches))) as executor:
            futures = {
                submit_with_context(executor, call_openai_vision_resilient, [batch_reference], our_product, same_product, payload.get("purpose", ""), start + 1): (reference_index, reference, start, len(batch_reference["image_urls"]))
                for reference_index, reference, start, batch_reference in batches
            }
            for future in as_completed(futures):
                reference_index, reference, start, batch_size = futures[future]
                try:
                    batch_result = future.result()
                except Exception as exc:
                    source_link = text(reference.get("url"), 900) or text(reference.get("asin"), 80) or "（链接缺失）"
                    source_asin = text(reference.get("asin"), 40)
                    identity = f"（ASIN {source_asin}）" if source_asin else ""
                    raise ValueError(
                        f"参考链接分析失败{identity}：{source_link}；"
                        f"图片第 {start + 1}-{start + batch_size} 张批次：{text(exc, 1100)}"
                    ) from None
                per_reference_batches[reference_index].append((start, batch_result))
                completed_batches += 1
                completed_images += len(batch_result.get("images", []))
                update_job(job_id, status="AI 正在归纳共性", progress=f"图片视觉识别已完成 {completed_batches}/{len(batches)} 个批次；正在整理类型共性", completed=completed_images, total=image_count)
        per_reference = []
        for reference, reference_batches in zip(references, per_reference_batches):
            reference_batches.sort(key=lambda entry: entry[0])
            analyses = [analysis for _, analysis in reference_batches]
            first = analyses[0] if analyses else {}
            merged = dict(first)
            merged["images"] = [image for analysis in analyses for image in analysis.get("images", [])]
            merged["core_benefits"] = list(dict.fromkeys(benefit for analysis in analyses for benefit in analysis.get("core_benefits", [])))
            merged["overall_structure_by_batch"] = [analysis.get("overall_structure", {}) for analysis in analyses if analysis.get("overall_structure")]
            per_reference.append(merged)
        update_job(job_id, status="正在汇总套图规律", progress=f"已识别 {image_count} 张图片并汇总 {len(references)} 条链接；正在生成四条共性总结和分类建议", completed=image_count, total=image_count)
        combined = summarize_reference_analyses(references, per_reference, our_product, same_product, payload.get("purpose", ""))
        combined["model"] = get_ai_config()[2] or "服务默认模型"
        analysis = combined
        update_job(job_id, status="分析完成", progress=f"已归纳 {len(references)} 条参考链接、{image_count} 张图片的共性，生成四条套图总结及分类建议。", result=analysis, done=True, completed=image_count, total=image_count)
    except Exception as exc:
        update_job(job_id, status="分析失败", error=redact_secrets(text(exc, 1500), ai_config[0]), done=True)
    finally:
        _request_ai_config.reset(context_token)


def excel_text(value):
    value = text(value, 8000)
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def set_text(cell, value, hyperlink=False):
    value = excel_text(value)
    cell.value = value
    cell.data_type = "s"
    if hyperlink and value.startswith(("https://", "http://")):
        cell.hyperlink = value
    cell.alignment = copy.copy(cell.alignment)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def parse_excel_date(value):
    value = text(value, 40)
    if not value:
        return None
    try:
        return dt.date.fromisoformat(value[:10])
    except ValueError:
        return value


def extend_detail_rows(ws, required_count):
    base_count = 10
    extra = max(0, required_count - base_count)
    if not extra:
        return
    merges = []
    for merged in list(ws.merged_cells.ranges):
        merges.append((merged.min_col, merged.min_row, merged.max_col, merged.max_row))
        ws.unmerge_cells(str(merged))
    heights = {idx: dim.height for idx, dim in ws.row_dimensions.items() if dim.height is not None}
    for idx in list(ws.row_dimensions):
        del ws.row_dimensions[idx]
    ws.insert_rows(24, extra)
    source_row = 23
    for row in range(24, 24 + extra):
        for col in range(1, 13):
            source = ws.cell(source_row, col)
            target = ws.cell(row, col)
            if source.has_style:
                target._style = copy.copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            target.alignment = copy.copy(source.alignment)
        ws.row_dimensions[row].height = heights.get(source_row, 55)
    for row, height in heights.items():
        new_row = row + extra if row >= 24 else row
        ws.row_dimensions[new_row].height = height
    for min_col, min_row, max_col, max_row in merges:
        if min_row >= 24:
            min_row += extra
            max_row += extra
        elif max_row >= 24:
            max_row += extra
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
    ws.print_area = f"A1:L{33 + extra}"


def add_research_sheet(wb, research):
    sheet_name = "竞品分析"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:L1")
    ws["A1"] = "商品来源与图片记录"
    ws["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="193A5A")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:L2")
    ws["A2"] = "本表记录直采或 Bright Data 回退实际读取到的商品页字段及来源链接；来源状态标明采集方式。图片观察和 AI 分析请查看「AI套图分析与生图指令」工作表。"
    ws["A2"].font = Font(name="Microsoft YaHei", size=10, color="526274", italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30
    headers = ["来源角色", "产品名称", "品牌", "平台", "商品链接", "价格", "返回图片数量", "图片序号", "图片链接", "来源状态", "商品页卖点/描述", "规格参数"]
    header_fill = PatternFill("solid", fgColor="DCEAF5")
    thin = Side(style="thin", color="CBD5DF")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(4, col, header)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="193A5A")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[4].height = 30
    out_row = 5
    for item in research.get("items", []):
        images = item.get("image_urls") or []
        if not images:
            values = [item.get("role"), item.get("title"), item.get("brand"), item.get("platform"), item.get("url"), item.get("price"), item.get("images_count"), "", "", item.get("status"), "；".join(item.get("features") or []) or item.get("description"), json.dumps(item.get("specifications") or [], ensure_ascii=False)]
            for col, value in enumerate(values, 1):
                if col == 7 and value not in (None, ""):
                    ws.cell(out_row, col).value = int(value)
                else:
                    set_text(ws.cell(out_row, col), value, hyperlink=(col in (5, 9)))
            out_row += 1
            continue
        manual = item.get("manual_observations") or {}
        for index, url in enumerate(images, 1):
            note = manual.get(str(index), {}) if isinstance(manual, dict) else {}
            values = [item.get("role"), item.get("title"), item.get("brand"), item.get("platform"), item.get("url"), item.get("price"), item.get("images_count"), index, url, item.get("status"), "；".join(item.get("features") or []) or item.get("description"), json.dumps(item.get("specifications") or [], ensure_ascii=False)]
            for col, value in enumerate(values, 1):
                if col in (7, 8) and value not in (None, ""):
                    ws.cell(out_row, col).value = int(value)
                else:
                    set_text(ws.cell(out_row, col), value, hyperlink=(col in (5, 9)))
            out_row += 1
    for col, width in enumerate([14, 38, 20, 14, 46, 14, 14, 10, 56, 28, 50, 54], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(5, out_row):
        ws.row_dimensions[row].height = 54
        for col in range(1, 13):
            cell = ws.cell(row, col)
            cell.font = Font(name="Microsoft YaHei", size=9, color="263445")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:L{max(4, out_row - 1)}"
    for warning in research.get("warnings", []):
        ws.cell(out_row + 1, 1, "采集提醒")
        ws.cell(out_row + 1, 1).font = Font(name="Microsoft YaHei", bold=True, color="A64322")
        ws.merge_cells(start_row=out_row + 1, start_column=2, end_row=out_row + 1, end_column=12)
        ws.cell(out_row + 1, 2, excel_text(f"{warning.get('source', '来源')}：{warning.get('message', '')}"))
        ws.cell(out_row + 1, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[out_row + 1].height = 36
        out_row += 1


def add_ai_analysis_sheet(wb, analysis):
    if not analysis:
        return
    sheet_name = "AI套图分析与生图指令"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    navy = "193A5A"
    pale = "DCEAF5"
    thin = Side(style="thin", color="CBD5DF")
    ws.merge_cells("A1:J1")
    ws["A1"] = "AI 套图共性总结、分类建议与画面文案"
    ws["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ref = analysis.get("reference_product") or {}
    refs = analysis.get("reference_products") or ([ref] if ref else [])
    our = analysis.get("our_product") or {}
    summary_rows = [
        (f"参考商品（{len(refs)} 条）", "；".join(item.get("title") or item.get("asin") or "" for item in refs), "\n".join(item.get("url") or "" for item in refs)),
        ("我方商品", our.get("title") or ("与参考商品相同" if analysis.get("same_product") else "未提供"), our.get("url") or ref.get("url") or ""),
        ("商品分析", analysis.get("product_summary_zh", ""), ""),
        ("核心卖点", "；".join(analysis.get("core_benefits") or []), ""),
        ("差异结论", analysis.get("differentiation_zh", ""), ""),
        ("来源依据", analysis.get("source_basis_zh", ""), ""),
    ]
    for row_idx, (label, value, link) in enumerate(summary_rows, 3):
        ws.cell(row_idx, 1, label).font = Font(name="Microsoft YaHei", size=10, bold=True, color=navy)
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=7)
        set_text(ws.cell(row_idx, 2), value)
        ws.merge_cells(start_row=row_idx, start_column=8, end_row=row_idx, end_column=10)
        set_text(ws.cell(row_idx, 8), link, hyperlink=(row_idx == 4))
        ws.row_dimensions[row_idx].height = 30 if row_idx < 5 else 42
    structure = analysis.get("overall_structure") or {}
    ws.merge_cells("A10:J10")
    ws["A10"] = "套图整体规律"
    ws["A10"].font = Font(name="Microsoft YaHei", size=12, bold=True, color=navy)
    structure_labels = [("顺序与结构", "sequence"), ("视觉风格", "style"), ("构图", "composition"), ("色彩", "color"), ("材质表现", "material_presentation"), ("文案安排", "copy_strategy")]
    for row_idx, (label, key) in enumerate(structure_labels, 11):
        ws.cell(row_idx, 1, label).font = Font(name="Microsoft YaHei", size=9, bold=True, color=navy)
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=10)
        set_text(ws.cell(row_idx, 2), structure.get(key, ""))
        ws.row_dimensions[row_idx].height = 32
    ws.merge_cells("A18:J18")
    ws["A18"] = f"四条{analysis.get('summary_scope') or '套图'}总结（{analysis.get('summary_image_count') or analysis.get('image_count') or 0} 张参考图）"
    ws["A18"].font = Font(name="Microsoft YaHei", size=12, bold=True, color=navy)
    for row_idx, point in enumerate((analysis.get("summary_points") or [])[:4], 19):
        ws.cell(row_idx, 1, f"共性总结 {row_idx - 18}").font = Font(name="Microsoft YaHei", size=9, bold=True, color=navy)
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=10)
        set_text(ws.cell(row_idx, 2), point)
        ws.row_dimensions[row_idx].height = 32
    ws.merge_cells("A24:J24")
    ws["A24"] = "按图片类型给出的设计建议（不含逐张分析）"
    ws["A24"].font = Font(name="Microsoft YaHei", size=12, bold=True, color=navy)
    headers = ["图片类型", "建议序号", "同类图片共性", "设计建议", "对应买家关注点", "画面文案 / 卖点文字", "English on-image copy", "中文生图指令", "English image prompt", "参考来源链接"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(25, col, header)
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color=navy)
        cell.fill = PatternFill("solid", fgColor=pale)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    row_idx = 26
    for group in analysis.get("image_groups") or []:
        for suggestion_index, suggestion in enumerate(group.get("suggestions") or [], 1):
            sources = group.get("reference_sources") or []
            example_urls = [item.get("image_url") for item in group.get("example_images") or [] if item.get("image_url")]
            source_notes = ("商品页：\n" + "\n".join(sources) if sources else "")
            if example_urls:
                source_notes += ("\n" if source_notes else "") + "参考例图：\n" + "\n".join(example_urls)
            values = [group.get("type"), f"方案 {suggestion_index}", group.get("shared_patterns"), suggestion.get("direction"), suggestion.get("buyer_focus_zh"), suggestion.get("copy_zh"), suggestion.get("copy_en"), suggestion.get("prompt_zh"), suggestion.get("prompt_en"), source_notes]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row_idx, col)
                set_text(cell, value, hyperlink=False)
                cell.font = Font(name="Microsoft YaHei", size=9, color="263445")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
            ws.row_dimensions[row_idx].height = 84
            row_idx += 1
    widths = [21, 13, 34, 42, 28, 30, 30, 48, 48, 42]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A26"
    ws.auto_filter.ref = f"A25:J{max(25, row_idx - 1)}"


def export_workbook(payload):
    if not TEMPLATE_PATH.exists():
        raise ValueError("找不到模板文件：templates/做图需求表模版-10.12.xlsx")
    rows = payload.get("rows")
    analysis = payload.get("analysis") or {}
    if not isinstance(rows, list) or (not rows and not analysis):
        raise ValueError("请先完成套图分析，或至少生成一条图片需求明细。")
    research = payload.get("research") or {}
    product = payload.get("product") or {}
    if not text(product.get("name") or research.get("product_name")):
        raise ValueError("请填写产品名称或先采集一个有效商品页。")
    due_date = parse_excel_date(product.get("due_date")) if text(product.get("due_date")) else None
    if due_date is not None and not isinstance(due_date, dt.date):
        raise ValueError("期望交付日期格式无效。")
    required_fields = ("position", "type", "size", "content", "buyer_focus", "reference", "style", "delivery")
    for index, row in enumerate(rows, 1):
        missing = [key for key in required_fields if not text(row.get(key))]
        if text(row.get("type")) == "待确认":
            missing.append("type")
        try:
            quantity = float(row.get("quantity"))
            if quantity < 1 or not quantity.is_integer():
                missing.append("quantity")
        except (TypeError, ValueError):
            missing.append("quantity")
        if missing:
            raise ValueError(f"第 {index} 条图需信息未完成，请补齐内容、买家关注点、参考对象和交付要求。")
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb["做图需求表"]
    extend_detail_rows(ws, len(rows))

    request_id = "DJ-" + dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    metadata = {
        "B5": request_id,
        "D5": product.get("name") or research.get("product_name"),
        "B6": product.get("sku") or product.get("asin"),
        "D6": "US",
        "B7": product.get("submitter"),
        "D7": dt.date.today(),
        "B8": product.get("priority") or "中",
        "D8": due_date,
        "B9": product.get("purpose") or "商品图片需求；商品来源、套图分析和图上文案见「竞品分析」与「AI套图分析与生图指令」工作表。",
        "B10": product.get("status") or "待提交",
    }
    for address, value in metadata.items():
        cell = ws[address]
        if value is None:
            value = ""
        if isinstance(value, (dt.date, dt.datetime)):
            cell.value = value
            cell.number_format = "yyyy-mm-dd"
        else:
            set_text(cell, value)
    field_map = [
        ("position", 2), ("type", 3), ("quantity", 4), ("size", 5),
        ("content", 6), ("buyer_focus", 7), ("copy", 8), ("reference", 9),
        ("style", 10), ("delivery", 11), ("notes", 12),
    ]
    for index, row in enumerate(rows):
        excel_row = 14 + index
        ws.cell(excel_row, 1).value = index + 1
        for key, col in field_map:
            cell = ws.cell(excel_row, col)
            value = row.get(key, "")
            if key == "quantity":
                cell.value = int(float(value))
            else:
                set_text(cell, value, hyperlink=(key == "reference"))
        ws.row_dimensions[excel_row].height = max(ws.row_dimensions[excel_row].height or 55, 55)

    add_research_sheet(wb, research)
    add_ai_analysis_sheet(wb, analysis)
    output = BytesIO()
    wb.save(output)
    data = output.getvalue()
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            bad_member = archive.testzip()
            if bad_member:
                raise ValueError(f"导出文件校验失败：{bad_member}")
    except zipfile.BadZipFile as exc:
        raise ValueError("生成的文件不是有效 Excel 工作簿。") from exc
    return data


class Handler(BaseHTTPRequestHandler):
    server_version = "ImageBriefLocal/1.0"

    def log_message(self, format_string, *args):
        print("[%s] %s" % (self.log_date_time_string(), format_string % args))

    def request_path(self):
        parsed = urllib.parse.urlsplit(self.path)
        if VERCEL_RUNTIME and parsed.path in {"/api/index.py", "/api/index"}:
            routed_path = urllib.parse.parse_qs(parsed.query).get("__vercel_path", [""])[0]
            if routed_path.startswith("/"):
                return urllib.parse.urlsplit(routed_path).path
        return parsed.path

    def send_json(self, value, status=200):
        encoded = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(encoded)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(encoded)

    def read_json(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            raise ValueError("请求长度无效。") from None
        if length <= 0 or length > MAX_BODY_BYTES:
            raise ValueError("请求内容为空或超过 12 MB。")
        try:
            return json.loads(self.rfile.read(length).decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise ValueError("请求内容不是有效 JSON。") from None

    def do_GET(self):
        path = self.request_path()
        if path == "/healthz":
            self.send_json({"ok": True, "template_ready": TEMPLATE_PATH.is_file()})
            return
        elif path == "/":
            target = WEB_DIR / "index.html"
        elif path in ("/style.css", "/app.js"):
            target = WEB_DIR / path.rsplit("/", 1)[-1]
        elif path == "/api/config":
            brightdata_token, brightdata_zone = get_brightdata_config()
            self.send_json({"api_base_url": _ai_api_base_url, "brightdata_configured": bool(brightdata_token and brightdata_zone), "brightdata_zone": brightdata_zone, "admin_required": bool(ADMIN_TOKEN), "deployment_config_readonly": VERCEL_RUNTIME, "collection_method": "direct_then_brightdata"})
            return
        elif path.startswith("/api/jobs/"):
            job = job_snapshot(path.rsplit("/", 1)[-1])
            if not job:
                self.send_json({"error": "采集任务不存在或服务已重启。"}, 404)
            else:
                self.send_json(job)
            return
        else:
            self.send_json({"error": "未找到请求的页面。"}, 404)
            return
        try:
            content = target.read_bytes()
        except OSError:
            self.send_json({"error": "页面资源暂不可用。"}, 404)
            return
        content_type = mimetypes.guess_type(target.name)[0] or "application/octet-stream"
        self.send_response(200)
        self.send_header("Content-Type", f"{content_type}; charset=utf-8")
        self.send_header("Content-Length", str(len(content)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(content)

    def do_POST(self):
        path = self.request_path()
        try:
            payload = self.read_json()
            if path == "/api/config":
                if VERCEL_RUNTIME:
                    self.send_json({"error": "Vercel 上的公共采集服务设置由站点管理员在 Vercel 环境变量中维护。"}, 409)
                    return
                if ADMIN_TOKEN and not hmac.compare_digest(text(payload.get("admin_token"), 500), ADMIN_TOKEN):
                    self.send_json({"error": "管理员密码不正确，不能修改公共服务设置。"}, 403)
                    return
                old_brightdata_token, old_brightdata_zone = get_brightdata_config()
                brightdata_token = text(payload.get("brightdata_api_token"), 1000).strip()
                brightdata_zone = text(payload.get("brightdata_zone", old_brightdata_zone), 200).strip()
                with _ai_config_lock:
                    global _brightdata_api_token, _brightdata_zone
                    _brightdata_api_token = brightdata_token if brightdata_token else ("" if payload.get("clear_brightdata_api_token") else old_brightdata_token)
                    _brightdata_zone = brightdata_zone
                    brightdata_configured = bool(_brightdata_api_token and _brightdata_zone)
                self.send_json({"brightdata_configured": brightdata_configured, "brightdata_zone": _brightdata_zone, "admin_required": bool(ADMIN_TOKEN), "message": "Bright Data 回退设置仅保存在当前服务进程内。"})
                return
            if path == "/api/models":
                api_key = text(self.headers.get("X-User-AI-Key"), 1000)
                if not api_key or any(ord(character) < 32 for character in api_key):
                    raise ValueError("请先填写你自己的 AI API Key。")
                base_url, _ = normalize_ai_api_base_url(payload.get("api_base_url", DEFAULT_AI_API_BASE_URL))
                if urllib.parse.urlsplit(base_url).scheme != "https":
                    raise ValueError("AI API Base URL 必须使用 HTTPS，避免传输 API Key 时被截取。")
                vision_model = "" if payload.get("use_provider_default") else text(payload.get("vision_model"), 200)
                context_token = _request_ai_config.set((api_key, base_url, vision_model))
                try:
                    self.send_json({"models": fetch_relay_models()})
                except ValueError as exc:
                    self.send_json({"error": redact_secrets(exc, api_key)}, 502)
                finally:
                    _request_ai_config.reset(context_token)
                return
            if path == "/api/research":
                references = payload.get("reference_urls") or []
                if isinstance(references, str):
                    references = [line.strip() for line in references.splitlines() if line.strip()]
                if not isinstance(references, list):
                    raise ValueError("参考商品链接请按每行一条的格式填写。")
                if len(references) > MAX_REFERENCE_LINKS:
                    raise ValueError(f"参考商品链接最多支持 {MAX_REFERENCE_LINKS} 个。")
                if not text(payload.get("product_url")) and not references:
                    raise ValueError("请填写我方商品链接或至少一个参考商品链接。")
                job_id = uuid.uuid4().hex
                job = {"job_id": job_id, "status": "排队中", "progress": "等待采集任务开始", "done": False, "completed": 0, "total": 0, "created_at": dt.datetime.now().isoformat(timespec="seconds")}
                with _jobs_lock:
                    _jobs[job_id] = job
                    while len(_jobs) > 80:
                        _jobs.pop(next(iter(_jobs)))
                if VERCEL_RUNTIME:
                    perform_research(job_id, payload)
                    completed_job = job_snapshot(job_id) or job
                    self.send_json(completed_job)
                else:
                    threading.Thread(target=perform_research, args=(job_id, payload), daemon=True).start()
                    self.send_json(job, 202)
                return
            if path == "/api/analyze":
                api_key = text(self.headers.get("X-User-AI-Key"), 1000)
                if not api_key or any(ord(character) < 32 for character in api_key):
                    raise ValueError("请先填写并保存你自己的 AI API Key。")
                base_url, _ = normalize_ai_api_base_url(payload.get("api_base_url", DEFAULT_AI_API_BASE_URL))
                if urllib.parse.urlsplit(base_url).scheme != "https":
                    raise ValueError("AI API Base URL 必须使用 HTTPS，避免传输 API Key 时被截取。")
                vision_model = "" if payload.get("use_provider_default") else text(payload.get("vision_model"), 200)
                ai_config = (api_key, base_url, vision_model)
                source_job_id = text(payload.get("research_job_id"), 80)
                if VERCEL_RUNTIME and isinstance(payload.get("research_result"), dict):
                    source_job = {"result": payload["research_result"], "done": True}
                else:
                    source_job = job_snapshot(source_job_id)
                if not source_job or not source_job.get("done") or not source_job.get("result"):
                    raise ValueError("商品采集结果已过期或尚未完成，请重新提交链接。")
                selected_product_ids = payload.get("selected_product_ids") or payload.get("selected_asins") or ([payload.get("selected_asin")] if payload.get("selected_asin") else [])
                if not selected_product_ids:
                    raise ValueError("请至少选择一条参考商品链接。")
                job_id = uuid.uuid4().hex
                job = {"job_id": job_id, "status": "排队中", "progress": "等待视觉分析", "done": False, "completed": 0, "total": 0, "created_at": dt.datetime.now().isoformat(timespec="seconds")}
                with _jobs_lock:
                    _jobs[job_id] = job
                    while len(_jobs) > 80:
                        _jobs.pop(next(iter(_jobs)))
                if VERCEL_RUNTIME:
                    perform_analysis(job_id, source_job, payload, ai_config)
                    completed_job = job_snapshot(job_id) or job
                    self.send_json(completed_job)
                else:
                    threading.Thread(target=perform_analysis, args=(job_id, source_job, payload, ai_config), daemon=True).start()
                    self.send_json(job, 202)
                return
            if path == "/api/export":
                data = export_workbook(payload)
                product = payload.get("product") or {}
                base = re.sub(r"[^A-Za-z0-9_-]+", "_", text(product.get("sku") or product.get("asin") or "design_brief", 80)).strip("_") or "design_brief"
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="image_brief_{base}.xlsx"')
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)
                return
            self.send_json({"error": "未找到请求的接口。"}, 404)
        except ValueError as exc:
            self.send_json({"error": str(exc)}, 400)
        except Exception as exc:
            self.send_json({"error": "服务器处理失败：" + text(exc, 1000)}, 500)


if __name__ == "__main__":
    if not TEMPLATE_PATH.exists():
        raise SystemExit(f"找不到模板：{TEMPLATE_PATH}")
    if HOST not in {"127.0.0.1", "localhost", "::1"} and len(ADMIN_TOKEN) < 24:
        raise SystemExit("公开网络绑定前必须设置至少 24 个字符的 IMAGEREQ_ADMIN_TOKEN，以保护公共服务设置。")
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"图需工作台已启动：http://{HOST}:{PORT}")
    print("商品页支持任意可公开访问的 HTTP(S) 网页；优先直采，失败时可使用 Bright Data 回退。按 Ctrl+C 停止。")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
